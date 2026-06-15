"""
OSINT Monitoring — Telegram a'zolarini yig'uvchi skaner (Telethon).

Asosiy xususiyat:
  * O'chirilgan hisoblar (Deleted Account) QO'SHIMCHA API SO'ROVISIZ aniqlanadi.
    `user.deleted` belgisi a'zolar ro'yxati bilan birga keladi, shuning uchun
    bunday hisoblar uchun GetFullUser (bio uchun) so'rovi YUBORILMAYDI — API tejaladi.

Ishlatish:
    1) pip install -r requirements.txt
    2) .env faylini to'ldiring (API_ID, API_HASH, namuna .env.example da)
    3) python scanner.py -e https://t.me/Durmon1643 -e @boshqa_guruh
"""

import argparse
import asyncio
import os
from datetime import datetime

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from telethon import TelegramClient
from telethon.errors import FloodWaitError
from telethon.tl.functions.users import GetFullUserRequest

load_dotenv()

API_ID = int(os.getenv("API_ID", "0"))
API_HASH = os.getenv("API_HASH", "")
SESSION = os.getenv("SESSION_NAME", "osint_session")
OUTPUT = os.getenv("OUTPUT_FILE", "OSINT_Monitoring.xlsx")

# Excel ustunlari (sizning fayldagi tartib bilan bir xil)
HEADERS = [
    "№", "Telegram ID", "Kanal ID", "Ism", "Familya", "Username",
    "Telefon", "Bio", "Ochiq Kanallar", "Maxfiy Kanal",
    "Manba Guruh/Kanal", "Qo'shilgan",
]


def is_deleted(user) -> bool:
    """O'chirilgan hisobni aniqlaydi — qo'shimcha so'rovsiz.

    1-daraja: server bergan rasmiy `deleted` belgisi.
    2-daraja: hamma maydon bo'sh bo'lsa (zaxira tekshiruv).
    Eslatma: ismdagi "Deleted Account" MATNIGA tayanmaymiz — u ishonchsiz.
    """
    if getattr(user, "deleted", False):
        return True
    return not (user.first_name or user.last_name or user.username or user.phone)


def open_workbook():
    """Mavjud Excelni ochadi yoki yangi sarlavhali fayl yaratadi."""
    if os.path.exists(OUTPUT):
        wb = load_workbook(OUTPUT)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "List1"
        ws.append(HEADERS)
    return wb, ws


async def scan_entity(client, ws, source: str, start_index: int) -> tuple[int, dict]:
    """Bitta guruh/kanal a'zolarini yig'adi. (yangi_index, statistika) qaytaradi."""
    stats = {"jami": 0, "ochirilgan": 0, "tirik": 0, "tejalган_sorov": 0}
    idx = start_index
    now = datetime.now().strftime("%Y.%m.%d %H:%M")

    entity = await client.get_entity(source)
    kanal_id = getattr(entity, "id", "")

    async for user in client.iter_participants(entity, aggressive=True):
        stats["jami"] += 1
        idx += 1

        # --- O'CHIRILGAN HISOB: GetFullUser YUBORMAYMIZ (so'rov tejaladi) ---
        if is_deleted(user):
            stats["ochirilgan"] += 1
            stats["tejalган_sorov"] += 1
            ws.append([
                idx, user.id, kanal_id, "Deleted Account", "", "",
                "", "", "", "", source, now,
            ])
            continue

        # --- TIRIK HISOB: bio uchun to'liq profilни so'raymiz ---
        bio = ""
        try:
            full = await client(GetFullUserRequest(user.id))
            bio = full.full_user.about or ""
        except FloodWaitError as e:
            print(f"  FloodWait: {e.seconds}s kutilmoqda...")
            await asyncio.sleep(e.seconds + 1)
        except Exception:
            pass

        stats["tirik"] += 1
        ws.append([
            idx, user.id, kanal_id,
            user.first_name or "", user.last_name or "",
            f"@{user.username}" if user.username else "",
            user.phone or "", bio, "", "", source, now,
        ])

    return idx, stats


async def main():
    parser = argparse.ArgumentParser(description="Telegram OSINT skaner")
    parser.add_argument("-e", "--entity", action="append", required=True,
                        help="Guruh/kanal (link, @username yoki ID). Bir nechta -e bo'lishi mumkin.")
    args = parser.parse_args()

    if not API_ID or not API_HASH:
        raise SystemExit("XATO: .env faylida API_ID va API_HASH ni to'ldiring.")

    wb, ws = open_workbook()
    start_index = ws.max_row - 1  # sarlavhani hisobga olmaymiz

    total = {"jami": 0, "ochirilgan": 0, "tirik": 0, "tejalган_sorov": 0}

    async with TelegramClient(SESSION, API_ID, API_HASH) as client:
        for source in args.entity:
            print(f"\n>>> Yig'ilmoqda: {source}")
            try:
                start_index, stats = await scan_entity(client, ws, source, start_index)
                for k in total:
                    total[k] += stats[k]
                print(f"    jami={stats['jami']} | tirik={stats['tirik']} | "
                      f"o'chirilgan={stats['ochirilgan']}")
            except Exception as e:
                print(f"    XATO ({source}): {e}")

    wb.save(OUTPUT)
    print(f"\n==================== YAKUN ====================")
    print(f"Jami a'zo:            {total['jami']}")
    print(f"Tirik hisob:          {total['tirik']}")
    print(f"O'chirilgan hisob:    {total['ochirilgan']}")
    print(f"Tejalgan API so'rovi: {total['tejalган_sorov']}  (GetFullUser yuborilmadi)")
    print(f"Saqlandi:             {OUTPUT}")


if __name__ == "__main__":
    asyncio.run(main())
